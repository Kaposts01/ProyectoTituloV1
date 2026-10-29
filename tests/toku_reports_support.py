"""Constructor de reportes .xlsx SINTETICOS de Toku para pruebas (TK-18). Cero datos reales.

Mismas columnas y orden que los reportes reales (incluida la columna duplicada «Medio de pago» en
inscripciones), titulo de 6 filas o encabezado en la fila 1, IDs ficticios y SENTINELAS de PII en las
columnas personales para comprobar que nunca llegan a la BD, a los logs ni a los errores.

Se importa como modulo plano (`from toku_reports_support import ...`, igual que `billing_requests_support`).
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

SENTINEL_NAME = "SENTINELA-NOMBRE"
SENTINEL_MAIL = "sentinela@example.invalid"
SENTINEL_PHONE = "SENTINELA-CELULAR"
SENTINEL_RUT = "00.000.001-K"
SENTINEL_TOKEN = "SENTINELA-TOKEN"
SENTINEL_CARD = "SENTINELA-TARJETA"
SENTINEL_MANDATE = "SENTINELA-MANDATO"
SENTINEL_DUP_PM = "SENTINELA-MEDIO-DUPLICADO"
SENTINEL_EXTRA = "SENTINELA-COLUMNA-EXTRA"
PII_SENTINELS = (
    SENTINEL_NAME, SENTINEL_MAIL, SENTINEL_PHONE, SENTINEL_RUT, SENTINEL_TOKEN, SENTINEL_CARD,
    SENTINEL_MANDATE, SENTINEL_DUP_PM, SENTINEL_EXTRA, "SENTINELA-EXTERNO", "SENTINELA-PAC",
)  # fmt: skip
PEPPER = "pepper-de-pruebas-1"


def product_text(n: int) -> str:
    """`ID Producto` ficticio: texto libre con un RUT FALSO (en el real contiene el RUT del socio)."""
    return f"Donación Mensual 10000 - {n:08d}-K - 01/01/2030"


def sub_id(n: int) -> str:
    return f"sub_T{n:04d}"


# Encabezados por tipo, en el orden de los reportes reales (con columnas personales intercaladas).
COLUMNS: dict[str, list[str]] = {
    "productos": [
        "ID", "Fecha de Creación", "Fecha de Eliminación", "ID Producto", "Cliente", "Nombre", "Mail", "Celular",
        "Rut", "ID Externo", "is_recurring", "Monto", "due_day", "Medio de Pago", "Número de Tarjeta", "token",
        "Bin", "Procesador de Pago", "Estado",
    ],
    "medios_pago": [
        "ID", "Fecha de Creación", "Fecha de Eliminación", "ID Producto", "Cliente", "Nombre", "Mail", "Rut",
        "Tipo de Tarjeta", "Marca Tarjeta", "Banco Emisor", "Número de Tarjeta", "token", "Bin", "Estado",
        "Procesador de Pago", "last_payment_intent_response_code",
        "last_payment_intent_original_response_message",
    ],
    "intentos": [
        "ID", "Fecha de Creación", "Fecha de Eliminación", "Cliente", "Nombre", "Email", "Rut", "Deuda",
        "ID Producto", "Fecha de Vencimiento", "Monto", "Estado", "Código de Respuesta Toku", "Procesador de Pago",
        "Fecha de Transacción", "ID Mandato", "Número de Cuenta", "Rut PAC", "Nombre PAC",
    ],
    "deudas": [
        "ID", "Fecha de Creación", "Fecha de Eliminación", "ID Producto", "Fecha de Vencimiento", "Monto", "Pagado",
        "Anulado", "Cliente", "Nombre", "Mail", "Rut", "Estado",
    ],
    "inscripciones": [
        "ID", "Fecha de Creación", "Fecha de Eliminación", "Cliente", "Nombre", "Mail", "Medio de pago",
        "Medio de pago#2", "Estado transacción", "Código de Respuesta", "Estado Medio de Pago",
        "Fecha Último Intento de Pago", "Estado Último Intento de Pago", "id_subscription",
    ],
}  # fmt: skip
TITLES = {
    "productos": "Productos", "medios_pago": "Medios de pago", "intentos": "Intentos de pago",
    "deudas": "Deudas", "inscripciones": "Inscripciones de medio de pago",
}  # fmt: skip
PII_VALUES = {
    "Nombre": SENTINEL_NAME, "Mail": SENTINEL_MAIL, "Email": SENTINEL_MAIL, "Celular": SENTINEL_PHONE,
    "Rut": SENTINEL_RUT, "ID Externo": "SENTINELA-EXTERNO", "Número de Tarjeta": SENTINEL_CARD,
    "token": SENTINEL_TOKEN, "Bin": "SENTINELA-BIN", "Número de Cuenta": SENTINEL_CARD,
    "Rut PAC": SENTINEL_RUT, "Nombre PAC": "SENTINELA-PAC",
}  # fmt: skip


def default_row(report_type: str, n: int = 1, **over: Any) -> dict[str, Any]:
    """Fila con valores por defecto y sentinelas de PII. `over` pisa por NOMBRE EXACTO de columna."""
    created = "2026-01-01T10:00:00"
    base: dict[str, Any]
    if report_type == "productos":
        base = {
            "ID": sub_id(n), "Fecha de Creación": created, "Fecha de Eliminación": None,
            "ID Producto": product_text(n), "Cliente": f"cus_T{n:04d}", "is_recurring": "true", "Monto": 10000,
            "due_day": 5, "Medio de Pago": f"pm_T{n:04d}", "Procesador de Pago": "Transbank", "Estado": "chargeable",
        }  # fmt: skip
    elif report_type == "medios_pago":
        base = {
            "ID": f"pm_T{n:04d}", "Fecha de Creación": created, "Fecha de Eliminación": None,
            "ID Producto": product_text(n), "Cliente": f"cus_T{n:04d}", "Tipo de Tarjeta": "credit",
            "Marca Tarjeta": "visa", "Banco Emisor": "Banco Ficticio", "Estado": "chargeable",
            "Procesador de Pago": "Transbank", "last_payment_intent_response_code": None,
            "last_payment_intent_original_response_message": None,
        }  # fmt: skip
    elif report_type == "intentos":
        base = {
            "ID": f"pi_T{n:04d}", "Fecha de Creación": created, "Fecha de Eliminación": None,
            "Cliente": f"cus_T{n:04d}", "Deuda": f"in_T{n:04d}", "ID Producto": product_text(n),
            "Fecha de Vencimiento": "2026-10-05", "Monto": 10000, "Estado": "FALLIDO",
            "Código de Respuesta Toku": "1001", "Procesador de Pago": "Transbank",
            "Fecha de Transacción": "2026-10-06T08:00:00", "ID Mandato": SENTINEL_MANDATE,
        }  # fmt: skip
    elif report_type == "deudas":
        base = {
            "ID": f"in_T{n:04d}", "Fecha de Creación": created, "Fecha de Eliminación": None,
            "ID Producto": product_text(n), "Fecha de Vencimiento": "05/10/2026", "Monto": 10000, "Pagado": "false",
            "Anulado": "false", "Cliente": f"cus_T{n:04d}", "Estado": "overdue",
        }  # fmt: skip
    elif report_type == "inscripciones":
        base = {
            "ID": f"pmii_T{n:04d}", "Fecha de Creación": created, "Fecha de Eliminación": None,
            "Cliente": f"cus_T{n:04d}", "Medio de pago": f"pm_T{n:04d}", "Medio de pago#2": SENTINEL_DUP_PM,
            "Estado transacción": "failed", "Código de Respuesta": "TEPSP1001", "Estado Medio de Pago": "chargeable",
            "Fecha Último Intento de Pago": "2026-10-06T08:00:00", "Estado Último Intento de Pago": "failed",
            "id_subscription": sub_id(n),
        }  # fmt: skip
    else:
        raise ValueError(report_type)
    base.update(over)
    return base


DEFAULT_CATALOG = [
    ("Código Toku", "Gravedad", "Mensaje", "Descripción"),
    ("0000", "SU", "Aprobada", "Transacción aprobada"),
    ("1001", "TE", "Fondos insuficientes", "El medio no tiene fondos"),
    ("1019", "PE", "Tarjeta bloqueada", "Tarjeta bloqueada por el emisor"),
    ("9999", "UE", "Error desconocido", None),
]


def write_report(
    path: Path,
    report_type: str,
    rows: list[dict[str, Any]] | None = None,
    *,
    with_title: bool = True,
    generated: datetime | str | None = "2026-10-20 10:00:00",
    window: tuple[str, str] = ("2026-01-01", "2026-10-20"),
    title: str | None = None,
    columns: list[str] | None = None,
    drop_columns: tuple[str, ...] = (),
    extra_columns: dict[str, Any] | None = None,
    catalog: list[tuple] | None = None,
    n_rows: int = 3,
) -> Path:
    """Escribe un .xlsx sintetico. `rows` = dicts por nombre de columna (ver `default_row`)."""
    cols = [c for c in (columns or COLUMNS[report_type]) if c not in drop_columns]
    if rows is None:
        rows = [default_row(report_type, i + 1) for i in range(n_rows)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    r = 1
    if with_title:
        meta = [
            f"Reporte | {title or TITLES[report_type]}",
            f"Fecha de generación: {generated}" if generated else "Fecha de generación",
            "Tipo de fecha: Fecha de creación",
            f"Fecha de inicio: {window[0]}",
            f"Fecha de término: {window[1]}",
            None,
        ]
        for text in meta:
            if text is not None:
                ws.cell(row=r, column=1, value=text)
            r += 1
    extras = list((extra_columns or {}).items())
    headers = [c.split("#")[0] for c in cols] + [name for name, _ in extras]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=r, column=i, value=h)
    r += 1
    for row in rows:
        merged = {**{k: v for k, v in PII_VALUES.items()}, **row}
        values = [merged.get(c) for c in cols] + [v for _, v in extras]
        for i, v in enumerate(values, start=1):
            if v is not None:
                ws.cell(row=r, column=i, value=v)
        r += 1
    if catalog is not None or report_type == "intentos":
        cs = wb.create_sheet("Códigos de Error Toku")
        for ri, line in enumerate(DEFAULT_CATALOG if catalog is None else catalog, start=1):
            for ci, v in enumerate(line, start=1):
                if v is not None:
                    cs.cell(row=ri, column=ci, value=v)
    wb.save(path)
    wb.close()
    return path


def write_unsupported(path: Path, name: str = "Clientes") -> Path:
    """Reporte de un tipo no soportado (columnas de «Clientes»; valores sentinela)."""
    wb = Workbook()
    ws = wb.active
    for i, h in enumerate(["ID", "Nombre", "Mail", "Rut", "Teléfono"], start=1):
        ws.cell(row=7, column=i, value=h)
        ws.cell(row=8, column=i, value=SENTINEL_NAME if i > 1 else "cus_T0001")
    ws.cell(row=1, column=1, value=f"Reporte | {name}")
    wb.save(path)
    wb.close()
    return path


def as_date(value: str) -> date:
    return date.fromisoformat(value)
